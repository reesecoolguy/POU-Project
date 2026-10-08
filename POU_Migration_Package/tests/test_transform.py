import csv
import json
from datetime import datetime
from pathlib import Path

import openpyxl
import pytest

from pou_tools.extract import extract
from pou_tools.transform import Config, transform, write_outputs

REAL = Path(__file__).resolve().parent.parent.parent / "POU_Inventory_Pilot Test _With_Badge.xlsm"

INV_H = ["Item ID", "Item Name", "Location", "Min Qty", "Max Qty", "On Hand Qty", "Barcode", "Area:", "Description",
         "Manufacturer", "Notes", "Priority", "Lead Time", "Cost per Unit"]


def build(tmp_path, inv=(), tx=(), audits=(), users=(("1", "A"),)):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inventory"
    ws.append(INV_H)
    for r in inv:
        r = list(r) + [None] * (len(INV_H) - len(r))
        ws.append(r)
    t = wb.create_sheet("Transactions"); t.append(["Timestamp", "Item ID", "Transaction Type", "Quantity", "User"])
    for r in tx: t.append(list(r))
    a = wb.create_sheet("InventoryAudit"); a.append(["Date/Time", "Item ID", "Previous Qty", "Counted Qty", "Variance", "User"])
    for r in audits: a.append(list(r))
    u = wb.create_sheet("Users"); u.append(["Badge ID", "Employee Name"])
    for r in users: u.append(list(r))
    p = tmp_path / "t.xlsx"
    wb.save(p)
    return transform(extract(p))


def row(item, loc, mn, mx, oh, name="Name", **kw):
    r = [item, name, loc, mn, mx, oh, None, "Area1", kw.get("desc"), kw.get("mfr"), kw.get("notes"), kw.get("prio"), kw.get("lead"), kw.get("cost")]
    return r


def cats(R, sev=None):
    return [e.category for e in R.exceptions if sev is None or e.severity == sev]


def test_ids_stay_text_and_leading_zeros_preserved(tmp_path):
    R = build(tmp_path, inv=[row("00123", "1-A", 1, 5, 2), row(100416, "1-A", 1, 5, 2), row("k1", "1-A", 1, 5, 2)],
              users=[("00045", "Zed"), (21904, "Num")])
    ids = [i["ItemID"] for i in R.items]
    assert "00123" in ids and "100416" in ids and "K1" in ids
    assert all(isinstance(i, str) for i in ids)
    emp = {e["BadgeID"] for e in R.employees}
    assert emp == {"00045", "21904"}
    assert "NUMERIC_ITEM_ID" in cats(R) and "NUMERIC_BADGE" in cats(R)


def test_repeated_item_id_at_two_locations_is_kept_as_two_stock_records(tmp_path):
    R = build(tmp_path, inv=[row("K1", "2-A", 12, 20, None, name="Fixture", desc=None, mfr="Layke"),
                             row("K1", "Blohm", 24, 48, 24, name="Fixture", desc="225-05161", mfr="Layke")])
    assert len(R.items) == 1 and len(R.stock) == 2
    assert {s["StockKey"] for s in R.stock} == {"K1|2-A", "K1|BLOHM"}
    assert R.items[0]["Description"] == "225-05161"          # single recorded value, flagged
    assert "MASTER_VALUE_FROM_SIBLING_ROW" in cats(R)
    assert "ITEM_AT_MULTIPLE_LOCATIONS" in cats(R)
    assert not R.held_inventory


def test_blank_onhand_is_never_zero_and_has_no_opening_request(tmp_path):
    R = build(tmp_path, inv=[row("A", "1-A", 1, 5, None), row("B", "1-A", 1, 5, 0)])
    assert [o["StockKey"] for o in R.opening_requests] == ["B|1-A"]
    assert R.opening_requests[0]["Quantity"] == 0               # real zero is preserved as a zero
    q = {r["StockKey"]: r for r in R.qty_recon}
    assert q["A|1-A"]["SourceOnHand"] == "" and q["A|1-A"]["Status"] == "NO_BALANCE_BLANK_IN_SOURCE"
    assert "BLANK_ONHAND" in cats(R, "WARN")
    assert all("OnHandQty" not in s for s in R.stock)           # balances never in the stock file


def test_onhand_above_max_is_not_capped(tmp_path):
    R = build(tmp_path, inv=[row("A", "1-A", 1, 4, 50)])
    assert R.opening_requests[0]["Quantity"] == 50
    assert "ONHAND_ABOVE_MAX" in cats(R, "WARN")


def test_duplicate_item_location_holds_all_rows(tmp_path):
    R = build(tmp_path, inv=[row("A", "1-A", 1, 4, 3), row("a ", "1-a", 1, 4, 9), row("B", "1-A", 1, 4, 1)])
    assert [s["StockKey"] for s in R.stock] == ["B|1-A"]
    assert len(R.held_inventory) == 2
    assert cats(R, "BLOCKER").count("DUPLICATE_ITEM_LOCATION") == 2
    assert not any(o["StockKey"].startswith("A|") for o in R.opening_requests)


def test_reserved_separator_and_blank_keys_are_held(tmp_path):
    R = build(tmp_path, inv=[row("A|B", "1-A", 1, 4, 3), row(None, "1-A", 1, 4, 3), row("C", None, 1, 4, 3), row("D", "1-A", 1, 4, 3)])
    assert [s["StockKey"] for s in R.stock] == ["D|1-A"]
    assert set(cats(R, "BLOCKER")) >= {"RESERVED_CHARACTER", "BLANK_ITEM_ID", "BLANK_LOCATION"}


def test_normalisation_is_reported_not_silent(tmp_path):
    R = build(tmp_path, inv=[row(" *k100* ", "1-a", 1, 4, 3)])
    assert R.stock[0]["ItemID"] == "K100"
    assert "ID_NORMALIZED" in cats(R)


def test_conflicting_master_values_are_blanked_and_listed(tmp_path):
    R = build(tmp_path, inv=[row("A", "1-A", 1, 4, 3, desc="one", mfr="X"), row("A", "1-B", 1, 4, 3, desc="two", mfr="X")])
    assert R.items[0]["Description"] == "" and R.items[0]["Manufacturer"] == "X"
    e = [x for x in R.exceptions if x.category == "MASTER_VALUE_CONFLICT"]
    assert e and "one" in e[0].value and "two" in e[0].value


def test_non_numeric_lead_time_is_kept_in_notes(tmp_path):
    R = build(tmp_path, inv=[row("A", "1-A", 1, 4, 3, lead="2 weeks", cost="$4.50", prio=1)])
    assert "2 weeks" in R.items[0]["Notes"]
    assert R.items[0]["UnitCost"] == "4.5" and R.items[0]["Priority"] == "1"
    assert "NON_NUMERIC_VALUE" in cats(R, "WARN")


def test_invalid_min_max_row_is_held(tmp_path):
    R = build(tmp_path, inv=[row("A", "1-A", "x", 4, 3), row("B", "1-A", 1, None, 3)])
    assert not R.stock and len(R.held_inventory) == 2


def test_legacy_ledger_provenance_and_no_invented_values(tmp_path):
    R = build(tmp_path,
              inv=[row("A", "1-A", 1, 9, 5), row("M", "1-A", 1, 9, 5), row("M", "1-B", 1, 9, 5)],
              tx=[(datetime(2026, 7, 1, 9, 42, 51), "A", "REMOVE", 2, "Reese McCord"),
                  (datetime(2026, 1, 15, 9, 0, 0), "A", "ADD", 5, "Reese McCord"),
                  (datetime(2026, 7, 1, 10, 0, 0), "M", "REMOVE", 1, "Bob"),
                  (datetime(2026, 7, 1, 10, 0, 0), "ZZZ", "ADD", 1, "Bob"),
                  (datetime(2026, 7, 1, 10, 0, 0), "A", "TRANSFER", 1, "Bob")],
              audits=[(datetime(2026, 7, 1, 11, 0, 0), "A", 4, 6, 99, "Reese McCord")])
    led = {l["LegacySource"]: l for l in R.legacy_ledger}
    t2 = led["Transactions!r2"]
    assert t2["LedgerType"] == "ISSUE" and t2["QtyDelta"] == -2 and t2["QtyBefore"] == "" and t2["QtyAfter"] == ""
    assert t2["Origin"] == "Legacy" and t2["AffectsBalance"] == "FALSE" and t2["PostingState"] == "Posted" and t2["SeqNo"] == ""
    assert t2["OccurredUtc"] == "2026-07-01T14:42:51Z"          # CDT = UTC-5
    assert led["Transactions!r3"]["OccurredUtc"] == "2026-01-15T15:00:00Z"   # CST = UTC-6
    assert t2["LegacyUser"] == "Reese McCord" and t2["EmployeeName"] == "" and t2["LegacyTimestampText"] == "2026-07-01T09:42:51"
    assert t2["StockKey"] == "A|1-A" and t2["LocationResolution"] == "UniqueAtCutover"
    m = led["Transactions!r4"]                                      # item at two locations: location NOT invented
    assert m["StockKey"] == "" and m["LocationResolution"] == "Unresolved" and m["ItemID"] == "M"
    assert led["Transactions!r5"]["StockKey"] == "" and "LEGACY_LOCATION_UNRESOLVED" in cats(R, "WARN")
    assert "Transactions!r6" not in led and any(h["Row"] == 6 for h in R.held_legacy)   # unknown type held
    a = led["InventoryAudit!r2"]
    assert a["QtyBefore"] == 4 and a["QtyAfter"] == 6 and a["QtyDelta"] == 2 and a["LedgerType"] == "AUDIT"
    assert "LEGACY_AUDIT_VARIANCE_MISMATCH" in cats(R)


def test_dst_ambiguous_time_flagged(tmp_path):
    R = build(tmp_path, inv=[row("A", "1-A", 1, 9, 5)], tx=[(datetime(2026, 11, 1, 1, 30, 0), "A", "ADD", 1, "x")])
    assert "AMBIGUOUS_LOCAL_TIME" in cats(R, "WARN")


def test_outputs_are_deterministic_and_headers_match_schema(tmp_path):
    R = build(tmp_path, inv=[row("A", "1-A", 1, 9, 5)], tx=[(datetime(2026, 7, 1, 9, 0, 0), "A", "ADD", 1, "x")])
    write_outputs(R, tmp_path / "o1"); write_outputs(R, tmp_path / "o2")
    for p in (tmp_path / "o1").rglob("*.csv"):
        assert p.read_bytes() == (tmp_path / "o2" / p.relative_to(tmp_path / "o1")).read_bytes()
    from pou_tools.import_data import STAGES, read_rows
    from pou_tools.schema import load_schema
    s = load_schema()
    for st, fname, ln, key in STAGES:
        read_rows(tmp_path / "o1" / "import" / fname, s[ln])    # raises if any column is not in the list schema


@pytest.mark.skipif(not REAL.exists(), reason="real workbook not present")
def test_real_workbook_golden_numbers():
    R = transform(extract(REAL))
    s = R.summary
    assert s["inventory_rows"] == 254 and s["stock_records_imported"] == 254 and s["stock_records_held"] == 0
    assert s["items"] == 251 and s["items_with_multiple_locations"] == 3 and s["locations"] == 22
    assert s["opening_requests"] == 252 and s["blank_or_invalid_onhand"] == 2 and s["sum_opening_quantity"] == 973
    assert s["legacy_ledger_rows"] == 26 and s["employees"] == 7
    assert s["low_stock_le_min"] == 71 and s["low_stock_lt_min"] == 9
    keys = {st["StockKey"] for st in R.stock}
    for k in ("K102516|2-A", "K102516|BLOHM", "K102517|2-A", "K102517|BLOHM", "K13471|2-D", "K13471|2-G"):
        assert k in keys
    assert all(r["Status"] == "OK" for r in R.row_counts)
    blanks = [q["StockKey"] for q in R.qty_recon if q["SourceOnHand"] == ""]
    assert sorted(blanks) == ["K102516|2-A", "K102517|2-A"]
    # numeric IDs preserved as text
    assert {"100416", "100674"} <= {i["ItemID"] for i in R.items}
