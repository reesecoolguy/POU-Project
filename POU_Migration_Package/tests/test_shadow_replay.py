"""Shadow/replay: workbook activity after a baseline snapshot is replayed into a TEST site and compared - without dual live posting."""
import openpyxl
import pytest

from conftest import SITE_OWNER
from test_flow_basic import env   # noqa: F401
from invariants import check
from pou_tools import shadow_replay as sr
from pou_tools.extract import extract
from pou_tools.import_data import run_all
from pou_tools.schema import load_schema
from pou_tools.transform import transform, write_outputs

INV_H = ["Item ID", "Item Name", "Location", "Min Qty", "Max Qty", "On Hand Qty", "Barcode", "Area:", "Description", "Manufacturer", "Notes", "Priority", "Lead Time", "Cost per Unit"]


def workbook(tmp_path, name, inv, tx, audits):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Inventory"; ws.append(INV_H)
    for r in inv:
        ws.append(list(r) + [None] * (len(INV_H) - len(r)))
    t = wb.create_sheet("Transactions"); t.append(["Timestamp", "Item ID", "Transaction Type", "Quantity", "User"])
    for r in tx: t.append(list(r))
    a = wb.create_sheet("InventoryAudit"); a.append(["Date/Time", "Item ID", "Previous Qty", "Counted Qty", "Variance", "User"])
    for r in audits: a.append(list(r))
    u = wb.create_sheet("Users"); u.append(["Badge ID", "Employee Name"]); u.append(["77", "Pat"])
    p = tmp_path / name; wb.save(p)
    return extract(p)


from datetime import datetime
D = lambda d, h=9: datetime(2026, 8, d, h, 0, 0)
BASE_INV = [("B", "Bolt", "1-A", 1, 20, 10), ("C", "Cap", "1-A", 1, 20, 8), ("D", "Dowel", "2-B", 1, 20, 6), ("D", "Dowel", "1-A", 1, 20, 4), ("E", "Eye", "2-B", 1, 20, 7)]


def test_shadow_replay_matches_explains_and_flags(env, tmp_path):
    baseline = workbook(tmp_path, "base.xlsx", BASE_INV, [(D(1), "B", "REMOVE", 1, "Pat")], [])
    # LATER the live workbook: more movements (and a count on C that overwrote On Hand)
    cur_inv = [("B", "Bolt", "1-A", 1, 20, 10 - 2 + 5), ("C", "Cap", "1-A", 1, 20, 11), ("D", "Dowel", "2-B", 1, 20, 6), ("D", "Dowel", "1-A", 1, 20, 4 - 1), ("E", "Eye", "2-B", 1, 20, 7)]
    cur_tx = [(D(1), "B", "REMOVE", 1, "Pat"), (D(2), "B", "REMOVE", 2, "Pat"), (D(3), "B", "ADD", 5, "Pat"),
              (D(3, 10), "D", "REMOVE", 1, "Pat"),                    # D is in two locations: cannot be attributed
              (D(4), "ZZ", "REMOVE", 1, "Pat")]                       # unknown item
    cur_aud = [(D(5), "C", 8, 11, 3, "Pat")]                          # a physical count on C: +3, overwrote On Hand
    current = workbook(tmp_path, "cur.xlsx", cur_inv, cur_tx, cur_aud)

    plan = sr.build_plan(baseline, current)
    st = {m.request_id: m.status for m in plan.moves}
    assert st == {"SHADOW-T-r3": "REPLAYABLE", "SHADOW-T-r4": "REPLAYABLE", "SHADOW-T-r5": "AMBIGUOUS_LOCATION", "SHADOW-T-r6": "UNKNOWN_ITEM"}
    assert len(plan.audits) == 1

    # the TEST site = production procedure on the baseline: import + post the OPENING balances through the flow
    out = tmp_path / "mig"; write_outputs(transform(baseline), out)
    owner = env.c(SITE_OWNER)
    env.setting("ServerSessionMaxIdleMinutes", "720")
    run_all(owner, load_schema(), out / "import", ["locations", "employees", "items", "stock", "openings"], apply=True, log=lambda *a: None)
    env.clock.advance(120)
    env.sweep(max_actions=1_000_000)
    assert env.stock("B|1-A")["OnHandQty"] == 10           # the baseline opening is the workbook figure

    n_req = env.fake.count("POURequests")
    dry = sr.replay(owner, plan, apply=False, log=lambda *a: None)
    assert dry["created"] == 0 and env.fake.count("POURequests") == n_req          # a dry run writes nothing
    res = sr.replay(owner, plan, apply=True, log=lambda *a: None)
    assert res["created"] == 2 and res["skipped"] == 2
    assert sr.replay(owner, plan, apply=True, log=lambda *a: None)["existing"] == 2        # idempotent
    rows = {r["StockKey"]: r for r in sr.compare(owner, current, plan)}
    assert rows["B|1-A"]["Status"] == "NOT_YET_PROCESSED"
    env.clock.advance(120)
    env.sweep(max_actions=1_000_000)
    rows = {r["StockKey"]: r for r in sr.compare(owner, current, plan)}
    assert rows["B|1-A"]["Status"] == "MATCH" and rows["B|1-A"]["TestSiteOnHand"] == 13
    assert rows["C|1-A"]["Status"] == "EXPLAINED_BY_AUDIT" and rows["C|1-A"]["Difference"] == -3
    assert rows["D|1-A"]["Status"] == "AMBIGUOUS_ITEM" and rows["D|2-B"]["Status"] == "AMBIGUOUS_ITEM"
    assert rows["E|2-B"]["Status"] == "MATCH"
    check(env)
